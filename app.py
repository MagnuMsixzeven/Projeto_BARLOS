import os, re, secrets, sqlite3, hashlib, json, base64, threading
from datetime import datetime, timedelta, date
from functools import wraps
from flask import (Flask, render_template, request, redirect, url_for,
                   session, flash, jsonify, g)

try:
    from pywebpush import webpush as _webpush, WebPushException
    from py_vapid import Vapid as _Vapid
    from cryptography.hazmat.primitives.serialization import (
        Encoding as _Enc, PublicFormat as _PubFmt,
        PrivateFormat as _PrivFmt, NoEncryption as _NoEnc
    )
    _PUSH_OK = True
except ImportError:
    _PUSH_OK = False

_VAPID_CLAIMS = {"sub": "mailto:barberbook@barbearia.com.br"}

app = Flask(__name__)
app.secret_key = os.environ.get('BB_SECRET', 'barberbook-dev-secret-2024')

DB_PATH = os.path.join(os.path.dirname(__file__), 'barberbook.db')

# ── helpers ────────────────────────────────────────────────────────────────────
def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(e=None):
    db = g.pop('db', None)
    if db: db.close()

def query(sql, args=(), one=False, commit=False):
    db = get_db()
    cur = db.execute(sql, args)
    if commit:
        db.commit()
        return cur.lastrowid
    rv = cur.fetchone() if one else cur.fetchall()
    return rv

def hash_pw(pw): return hashlib.sha256(pw.encode()).hexdigest()

def login_required(f):
    @wraps(f)
    def deco(*a, **kw):
        if 'uid' not in session:
            flash('Faça login para continuar.', 'warning')
            return redirect(url_for('login'))
        return f(*a, **kw)
    return deco

def admin_required(f):
    @wraps(f)
    def deco(*a, **kw):
        if session.get('role') not in ('admin', 'gerente'):
            flash('Acesso restrito.', 'danger')
            return redirect(url_for('dashboard'))
        return f(*a, **kw)
    return deco

def fmt_brl(v):
    try: return f"{float(v):,.2f}".replace(',','X').replace('.',',').replace('X','.')
    except: return '0,00'

def _cfg():
    return {r['chave']: r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}

def _normalizar_tel(telefone):
    dig = ''.join(ch for ch in str(telefone or '') if ch.isdigit())
    if not dig:
        return ''
    return dig if dig.startswith('55') else ('55' + dig if len(dig) >= 10 else dig)

def _send_whatsapp_message(telefone, mensagem):
    """Envia mensagem pelo WhatsApp Cloud API, se configurado."""
    import urllib.request
    cfg = _cfg()
    phone_id = (cfg.get('whatsapp_phone_id') or '').strip()
    token = (cfg.get('whatsapp_api_token') or '').strip()
    tel = _normalizar_tel(telefone)
    if not tel or not phone_id or not token or not mensagem:
        return {'ok': False, 'msg': 'Integração WhatsApp não configurada.'}
    payload = json.dumps({
        'messaging_product': 'whatsapp',
        'to': tel,
        'type': 'text',
        'text': {'body': mensagem}
    }).encode('utf-8')
    req = urllib.request.Request(
        f'https://graph.facebook.com/v21.0/{phone_id}/messages',
        data=payload,
        method='POST',
        headers={
            'Authorization': f'Bearer {token}',
            'Content-Type': 'application/json'
        }
    )
    try:
        with urllib.request.urlopen(req, timeout=15) as resp:
            return {'ok': True, 'msg': resp.read().decode('utf-8', 'ignore')}
    except Exception as exc:
        return {'ok': False, 'msg': str(exc)}

def _parse_whatsapp_agendamento(texto):
    """Extrai data/hora/barbeiro/serviço de uma mensagem livre do WhatsApp."""
    texto = (texto or '').strip()
    low = texto.lower()
    info = {
        'data': '', 'hora': '',
        'barbeiro_id': None, 'barbeiro_nome': '',
        'servico_id': None, 'servico_nome': ''
    }

    servicos = [dict(r) for r in query("SELECT id,nome FROM servicos WHERE ativo=1 ORDER BY LENGTH(nome) DESC")]
    for s in servicos:
        nome_low = s['nome'].lower()
        if nome_low in low or nome_low.split()[0] in low:
            info['servico_id'] = s['id']
            info['servico_nome'] = s['nome']
            break
    if not info['servico_id']:
        if 'corte' in low and 'barba' in low:
            ach = next((s for s in servicos if 'corte' in s['nome'].lower() and 'barba' in s['nome'].lower()), None)
        elif 'barba' in low:
            ach = next((s for s in servicos if 'barba' in s['nome'].lower()), None)
        elif 'sobrancelha' in low:
            ach = next((s for s in servicos if 'sobrancelha' in s['nome'].lower()), None)
        else:
            ach = next((s for s in servicos if 'corte' in s['nome'].lower()), None)
        if ach:
            info['servico_id'] = ach['id']
            info['servico_nome'] = ach['nome']

    barbeiros = [dict(r) for r in query("SELECT id,nome FROM barbeiros WHERE ativo=1 ORDER BY nome")]
    for b in barbeiros:
        nome_low = b['nome'].lower()
        primeiro = nome_low.split()[0]
        if nome_low in low or primeiro in low:
            info['barbeiro_id'] = b['id']
            info['barbeiro_nome'] = b['nome']
            break
    if not info['barbeiro_id'] and barbeiros:
        info['barbeiro_id'] = barbeiros[0]['id']
        info['barbeiro_nome'] = barbeiros[0]['nome']

    m_data = re.search(r'\b(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?\b', low)
    if m_data:
        dia = int(m_data.group(1))
        mes = int(m_data.group(2))
        ano = int(m_data.group(3) or date.today().year)
        if ano < 100:
            ano += 2000
        try:
            info['data'] = date(ano, mes, dia).isoformat()
        except Exception:
            pass
    elif 'amanhã' in low or 'amanha' in low:
        info['data'] = (date.today() + timedelta(days=1)).isoformat()
    elif 'hoje' in low:
        info['data'] = date.today().isoformat()
    else:
        dias_map = {
            'segunda': 0, 'segunda-feira': 0,
            'terça': 1, 'terca': 1, 'terça-feira': 1, 'terca-feira': 1,
            'quarta': 2, 'quarta-feira': 2,
            'quinta': 3, 'quinta-feira': 3,
            'sexta': 4, 'sexta-feira': 4,
            'sábado': 5, 'sabado': 5,
            'domingo': 6
        }
        hoje_wd = date.today().weekday()
        for nome_dia, alvo in dias_map.items():
            if nome_dia in low:
                delta = (alvo - hoje_wd) % 7
                delta = 7 if delta == 0 else delta
                info['data'] = (date.today() + timedelta(days=delta)).isoformat()
                break

    m_hora = re.search(r'\b([01]?\d|2[0-3])(?:[:h](\d{2}))\b', low)
    if not m_hora:
        m_hora = re.search(r'\b([01]?\d|2[0-3])h\b', low)
    if m_hora:
        hh = int(m_hora.group(1))
        mm = int(m_hora.group(2) or 0) if m_hora.lastindex and m_hora.lastindex >= 2 and m_hora.group(2) else 0
        info['hora'] = f'{hh:02d}:{mm:02d}'

    return info

def _registrar_whatsapp_entrada(nome, telefone, mensagem, origem='webhook', payload=None):
    """Salva mensagem recebida e cria solicitação de agendamento quando possível."""
    tel = _normalizar_tel(telefone)
    info = _parse_whatsapp_agendamento(mensagem)
    status_parse = 'Mensagem recebida'
    agendamento_id = None

    dados_completos = all([tel, info['data'], info['hora'], info['servico_id'], info['barbeiro_id']])
    if dados_completos:
        indisponivel = query(
            "SELECT id FROM agendamentos WHERE data=? AND hora=? AND barbeiro_id=? AND status NOT IN ('Cancelado')",
            (info['data'], info['hora'], int(info['barbeiro_id']))
        )
        if indisponivel:
            status_parse = 'Horário solicitado já está ocupado'
        else:
            cli = query("SELECT id FROM clientes WHERE REPLACE(REPLACE(REPLACE(REPLACE(telefone,'(',''),')',''),' ',''),'-','') LIKE ?",
                        ('%'+tel[-11:],), one=True)
            if cli:
                cli_id = cli['id']
                query("UPDATE clientes SET nome=? WHERE id=?", (nome or 'Cliente WhatsApp', cli_id), commit=True)
            else:
                cli_id = query("INSERT INTO clientes(nome,telefone) VALUES(?,?)", (nome or 'Cliente WhatsApp', telefone), commit=True)
            agendamento_id = query(
                "INSERT INTO agendamentos(cliente_id,barbeiro_id,servico_id,data,hora,status,obs) VALUES(?,?,?,?,?,'Solicitado via WhatsApp',?)",
                (cli_id, int(info['barbeiro_id']), int(info['servico_id']), info['data'], info['hora'], 'Solicitação recebida automaticamente pelo WhatsApp.'),
                commit=True
            )
            _notify_barber_async(
                int(info['barbeiro_id']),
                'Novo pedido via WhatsApp',
                f"{(nome or 'Cliente').split()[0]} · {info['servico_nome'] or 'Serviço'} · {info['hora']}"
            )
            status_parse = f'Agendamento lançado no sistema #{agendamento_id}'

    query(
        """INSERT INTO whatsapp_mensagens(nome,telefone,direcao,mensagem,origem,data_solicitada,hora_solicitada,
              barbeiro_solicitado,servico_solicitado,status_parse,agendamento_id,lida,payload_json)
           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (
            nome or 'Cliente', telefone or tel, 'entrada', mensagem or '', origem,
            info['data'] or '', info['hora'] or '', info['barbeiro_nome'] or '', info['servico_nome'] or '',
            status_parse, agendamento_id, 0, json.dumps(payload or {}, ensure_ascii=False)
        ),
        commit=True
    )
    return {'ok': True, 'agendamento_id': agendamento_id, 'status_parse': status_parse, 'info': info}

@app.context_processor
def inject_whatsapp_badge():
    if 'uid' not in session:
        return {'wpp_unread_count': 0}
    try:
        unread = query("SELECT COUNT(*) c FROM whatsapp_mensagens WHERE lida=0", one=True)['c']
        pend = query("SELECT COUNT(*) c FROM agendamentos WHERE status='Solicitado via WhatsApp'", one=True)['c']
        return {'wpp_unread_count': int(unread or 0) + int(pend or 0)}
    except Exception:
        return {'wpp_unread_count': 0}

# ── init db ───────────────────────────────────────────────────────────────────
def init_db():
    db = sqlite3.connect(DB_PATH)
    db.executescript("""
        CREATE TABLE IF NOT EXISTS usuarios (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            usuario TEXT UNIQUE NOT NULL,
            senha TEXT NOT NULL,
            role TEXT DEFAULT 'barbeiro',
            ativo INTEGER DEFAULT 1,
            criado_em TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS barbeiros (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            telefone TEXT,
            especialidade TEXT,
            usuario_id INTEGER REFERENCES usuarios(id),
            ativo INTEGER DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS clientes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            telefone TEXT,
            email TEXT,
            bairro TEXT,
            criado_em TEXT DEFAULT (date('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS servicos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT NOT NULL,
            descricao TEXT,
            preco REAL DEFAULT 0,
            duracao INTEGER DEFAULT 30,
            ativo INTEGER DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS agendamentos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cliente_id INTEGER REFERENCES clientes(id) ON DELETE SET NULL,
            barbeiro_id INTEGER REFERENCES barbeiros(id) ON DELETE SET NULL,
            servico_id INTEGER REFERENCES servicos(id) ON DELETE SET NULL,
            data TEXT NOT NULL,
            hora TEXT NOT NULL,
            status TEXT DEFAULT 'Pendente',
            obs TEXT,
            criado_em TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS configuracoes (
            chave TEXT PRIMARY KEY,
            valor TEXT
        );
        CREATE TABLE IF NOT EXISTS horarios (
            dia TEXT PRIMARY KEY,
            ativo INTEGER DEFAULT 1,
            inicio TEXT DEFAULT '08:00',
            fim TEXT DEFAULT '18:00'
        );
        CREATE TABLE IF NOT EXISTS whatsapp_mensagens (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nome TEXT,
            telefone TEXT NOT NULL,
            direcao TEXT DEFAULT 'entrada',
            mensagem TEXT NOT NULL,
            origem TEXT DEFAULT 'webhook',
            data_solicitada TEXT DEFAULT '',
            hora_solicitada TEXT DEFAULT '',
            barbeiro_solicitado TEXT DEFAULT '',
            servico_solicitado TEXT DEFAULT '',
            status_parse TEXT DEFAULT '',
            agendamento_id INTEGER REFERENCES agendamentos(id) ON DELETE SET NULL,
            lida INTEGER DEFAULT 0,
            payload_json TEXT DEFAULT '',
            criado_em TEXT DEFAULT (datetime('now','localtime'))
        );
    """)
    # migrations – adiciona colunas se não existirem
    for _mig in [
        "ALTER TABLE barbeiros ADD COLUMN em_atendimento INTEGER DEFAULT 0",
        "ALTER TABLE barbeiros ADD COLUMN email TEXT DEFAULT ''",
        "ALTER TABLE agendamentos ADD COLUMN pagamento TEXT DEFAULT ''",
        "ALTER TABLE usuarios ADD COLUMN email TEXT DEFAULT ''",
        "ALTER TABLE usuarios ADD COLUMN foto TEXT DEFAULT ''",
    ]:
        try: db.execute(_mig)
        except: pass
    # push_subscriptions
    db.execute("""CREATE TABLE IF NOT EXISTS push_subscriptions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        barbeiro_id INTEGER NOT NULL,
        subscription_json TEXT NOT NULL UNIQUE,
        criado_em TEXT DEFAULT (datetime('now','localtime'))
    )""")
    # VAPID keys auto-gerados na primeira execução
    if _PUSH_OK:
        if not db.execute("SELECT valor FROM configuracoes WHERE chave='vapid_public'").fetchone():
            _v = _Vapid()
            _v.generate_keys()
            _pub = base64.urlsafe_b64encode(
                _v.public_key.public_bytes(_Enc.X962, _PubFmt.UncompressedPoint)
            ).rstrip(b'=').decode()
            _priv = _v.private_key.private_bytes(_Enc.PEM, _PrivFmt.TraditionalOpenSSL, _NoEnc()).decode()
            db.execute("INSERT OR REPLACE INTO configuracoes(chave,valor) VALUES('vapid_public',?)", (_pub,))
            db.execute("INSERT OR REPLACE INTO configuracoes(chave,valor) VALUES('vapid_private',?)", (_priv,))
    # seed admin
    adm = db.execute("SELECT id FROM usuarios WHERE usuario='admin'").fetchone()
    if not adm:
        db.execute("INSERT INTO usuarios(nome,usuario,senha,role) VALUES(?,?,?,?)",
                   ('Administrador','admin',hash_pw('admin123'),'admin'))
        db.execute("INSERT INTO usuarios(nome,usuario,senha,role) VALUES(?,?,?,?)",
                   ('Gerente','gerente',hash_pw('gerente123'),'gerente'))
        db.execute("INSERT INTO usuarios(nome,usuario,senha,role) VALUES(?,?,?,?)",
                   ('João Barbeiro','joao',hash_pw('joao123'),'barbeiro'))
    # seed barbeiros
    if not db.execute("SELECT id FROM barbeiros").fetchone():
        db.execute("INSERT INTO barbeiros(nome,telefone,especialidade,ativo) VALUES(?,?,?,1)",
                   ('João Silva','(11)99999-0001','Corte masculino'))
        db.execute("INSERT INTO barbeiros(nome,telefone,especialidade,ativo) VALUES(?,?,?,1)",
                   ('Pedro Costa','(11)99999-0002','Barba e bigode'))
        db.execute("INSERT INTO barbeiros(nome,telefone,especialidade,ativo) VALUES(?,?,?,1)",
                   ('Carlos Lima','(11)99999-0003','Corte + Barba'))
    # link barbeiro João ao usuário joao
    _ju = db.execute("SELECT id FROM usuarios WHERE usuario='joao'").fetchone()
    if _ju:
        db.execute("UPDATE barbeiros SET usuario_id=? WHERE nome LIKE 'Jo%o%' AND (usuario_id IS NULL OR usuario_id=0)",(_ju[0],))
    # seed servicos
    if not db.execute("SELECT id FROM servicos").fetchone():
        for n,p,d in [('Corte Masculino',35,30),('Barba',25,20),('Corte + Barba',55,50),('Combo Completo',70,60),('Sobrancelha',15,10)]:
            db.execute("INSERT INTO servicos(nome,preco,duracao) VALUES(?,?,?)",(n,p,d))
    # seed clientes
    if not db.execute("SELECT id FROM clientes").fetchone():
        for n,t in [('Lucas Oliveira','(11)91111-2222'),('Rafael Souza','(11)93333-4444'),('Marcos Pereira','(11)95555-6666'),('Gabriel Alves','(11)97777-8888')]:
            db.execute("INSERT INTO clientes(nome,telefone) VALUES(?,?)",(n,t))
    # seed config
    defaults = [('nome','Barbearia do Rei'),('telefone','(11)99999-0000'),('email','contato@barbeariadorei.com.br'),
                ('endereco','Rua das Flores, 123 – Centro'),('instagram','barbeariadorei'),
                ('link_agendamento','https://barberbook.app/barbeariadorei'),
                ('whatsapp_numero','5511999990000'),('whatsapp_phone_id',''),('whatsapp_api_token',''),
                ('whatsapp_verify_token', secrets.token_hex(12)),
                ('vencimento',(date.today()+timedelta(days=30)).strftime('%d/%m/%Y')),
                ('vencimento_iso',(date.today()+timedelta(days=30)).isoformat())]
    for k,v in defaults:
        db.execute("INSERT OR IGNORE INTO configuracoes(chave,valor) VALUES(?,?)",(k,v))
    # seed horarios
    dias = [('seg','Segunda',1),('ter','Terça',1),('qua','Quarta',1),('qui','Quinta',1),('sex','Sexta',1),('sab','Sábado',1),('dom','Domingo',0)]
    for k,_,a in dias:
        db.execute("INSERT OR IGNORE INTO horarios(dia,ativo) VALUES(?,?)",(k,a))
    # seed agendamentos
    if not db.execute("SELECT id FROM agendamentos").fetchone():
        hoje = date.today().isoformat()
        amanha = (date.today()+timedelta(1)).isoformat()
        for ci,bi,si,d,h,st in [
            (1,1,1,hoje,'09:00','Confirmado'),(2,1,2,hoje,'10:00','Pendente'),
            (3,2,3,hoje,'11:00','Confirmado'),(1,2,1,amanha,'09:30','Pendente'),
            (4,3,4,amanha,'14:00','Confirmado'),
        ]:
            db.execute("INSERT INTO agendamentos(cliente_id,barbeiro_id,servico_id,data,hora,status) VALUES(?,?,?,?,?,?)",
                       (ci,bi,si,d,h,st))
    db.commit()
    db.close()

# ── auth ──────────────────────────────────────────────────────────────────────
@app.route('/login', methods=['GET','POST'])
def login():
    if 'uid' in session: return redirect(url_for('dashboard'))
    error = None
    if request.method == 'POST':
        u = request.form.get('usuario','').strip()
        p = request.form.get('senha','')
        row = query("SELECT * FROM usuarios WHERE usuario=? AND ativo=1",(u,), one=True)
        if row and row['senha'] == hash_pw(p):
            session.update({'uid':row['id'],'nome':row['nome'],'usuario':row['usuario'],'role':row['role'],'email':row['email'] or ''})
            flash(f"Bem-vindo, {row['nome'].split()[0]}!", 'success')
            if row['role'] == 'barbeiro':
                return redirect(url_for('barbeiro_painel'))
            return redirect(url_for('dashboard'))
        error = 'Usuário ou senha incorretos.'
    return render_template('login.html', error=error)

@app.route('/logout')
def logout():
    session.clear()
    flash('Você saiu do sistema.', 'info')
    return redirect(url_for('login'))

@app.route('/demo')
def demo():
    session.clear()
    u = query("SELECT * FROM usuarios WHERE role IN ('admin','gerente') AND ativo=1 ORDER BY id LIMIT 1", one=True)
    if not u:
        return redirect(url_for('login'))
    session.update({'uid': u['id'], 'nome': u['nome'], 'usuario': u['usuario'], 'role': u['role'], 'email': u['email'] or ''})
    session['demo_mode'] = True
    return redirect(url_for('dashboard'))

# ── dashboard ─────────────────────────────────────────────────────────────────
@app.route('/dashboard')
@login_required
def dashboard():
    hoje = date.today().isoformat()
    mes  = date.today().strftime('%Y-%m')
    cfg  = {r['chave']:r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    # stats
    ag_hoje = query("SELECT COUNT(*) c FROM agendamentos WHERE data=?",(hoje,), one=True)['c']
    ag_mes  = query("SELECT COUNT(*) c FROM agendamentos WHERE data LIKE ?",(mes+'%',), one=True)['c']
    ag_pend = query("SELECT COUNT(*) c FROM agendamentos WHERE status IN ('Pendente','Solicitado via WhatsApp')",(), one=True)['c']
    t_cli   = query("SELECT COUNT(*) c FROM clientes",(), one=True)['c']
    novos   = query("SELECT COUNT(*) c FROM clientes WHERE criado_em LIKE ?",(mes+'%',), one=True)['c']
    t_barb  = query("SELECT COUNT(*) c FROM barbeiros WHERE ativo=1",(), one=True)['c']
    receita = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.data LIKE ? AND a.status='Concluido'",(mes+'%',), one=True)['v']
    # proximos hoje
    proximos_raw = query("""
        SELECT a.hora, c.nome cliente, s.nome servico, a.status
        FROM agendamentos a
        JOIN clientes c ON c.id=a.cliente_id
        JOIN servicos s ON s.id=a.servico_id
        WHERE a.data=? ORDER BY a.hora LIMIT 6
    """,(hoje,))
    proximos = [dict(r) for r in proximos_raw]
    # semana (7 dias a partir de hoje)
    semana = []
    nomes_dia = ['Dom','Seg','Ter','Qua','Qui','Sex','Sáb']
    for i in range(7):
        d = date.today() + timedelta(i)
        total = query("SELECT COUNT(*) c FROM agendamentos WHERE data=?",(d.isoformat(),), one=True)['c']
        semana.append({'idx':i,'nome':nomes_dia[d.weekday()%7],'num':d.day,'total':str(total) if total else '·','hoje':d==date.today(),'data':d.isoformat()})
    # contratos (placeholder)
    contratos = [{'nome':cfg.get('nome','Barbearia'),'plano':'PRO Mensal','inicio':'01/01/2025','status':'Ativo'}]
    return render_template('admin_dashboard.html',
        agora=datetime.now().strftime('%d/%m/%Y %H:%M'),
        barbearia_nome=cfg.get('nome',''), vencimento=cfg.get('vencimento',''),
        vencimento_iso=cfg.get('vencimento_iso',''), link_agendamento=cfg.get('link_agendamento',''),
        insta_handle=cfg.get('instagram',''), configuracoes_insta='#',
        agendamentos_hoje=ag_hoje, agendamentos_mes=ag_mes,
        agendamentos_pendentes=ag_pend, total_clientes=t_cli,
        novos_clientes=novos, total_barbeiros=t_barb,
        receita_mes=fmt_brl(receita), proximos=proximos,
        semana=semana, contratos=contratos)

# ── agendamentos ──────────────────────────────────────────────────────────────
@app.route('/agendamentos')
@login_required
def agendamentos():
    fd = request.args.get('data',''); fb = request.args.get('barbeiro',''); fs = request.args.get('status','')
    sql = """SELECT a.id, c.nome cliente, c.telefone telefone, s.nome servico, b.nome barbeiro,
                    a.data, a.hora, a.status, a.obs, a.servico_id, a.barbeiro_id, a.cliente_id, s.preco valor
             FROM agendamentos a
             JOIN clientes c ON c.id=a.cliente_id
             JOIN servicos s ON s.id=a.servico_id
             JOIN barbeiros b ON b.id=a.barbeiro_id
             WHERE 1=1"""
    args = []
    if fd: sql += " AND a.data=?"; args.append(fd)
    if fb: sql += " AND a.barbeiro_id=?"; args.append(fb)
    if fs: sql += " AND a.status=?"; args.append(fs)
    sql += " ORDER BY a.data DESC, a.hora"
    rows = [dict(r) for r in query(sql, args)]
    for r in rows:
        r['valor'] = fmt_brl(r['valor'])
        r['whatsapp_link'] = _wpp_link(r['id'], 'Olá {nome}! Seu horário para *{servico}* está registrado para {data} às {hora} com {barbeiro}. Responda aqui para confirmar')
    barbeiros_lista = [dict(r) for r in query("SELECT id,nome FROM barbeiros WHERE ativo=1 ORDER BY nome")]
    return render_template('admin_agendamentos.html', agendamentos=rows, total=len(rows),
                           barbeiros_lista=barbeiros_lista,
                           filtro_data=fd, filtro_barbeiro=fb, filtro_status=fs)

@app.route('/agendamentos/novo', methods=['GET','POST'])
@login_required
def novo_agendamento():
    if request.method == 'POST':
        f = request.form
        ag_id = query("INSERT INTO agendamentos(cliente_id,barbeiro_id,servico_id,data,hora,status,obs) VALUES(?,?,?,?,?,?,?)",
              (f['cliente_id'],f['barbeiro_id'],f['servico_id'],f['data'],f['hora'],f.get('status','Pendente'),f.get('obs','')), commit=True)
        _notificar_cliente_agendamento(ag_id, 'Olá {nome}! Seu agendamento de *{servico}* foi registrado para {data} às {hora} com {barbeiro}. Responda aqui para confirmar o horário')
        flash('Agendamento criado!', 'success')
        return redirect(url_for('agendamentos'))
    return render_template('form_agendamento.html',
        clientes=query("SELECT id,nome FROM clientes ORDER BY nome"),
        barbeiros=query("SELECT id,nome FROM barbeiros WHERE ativo=1 ORDER BY nome"),
        servicos=query("SELECT id,nome,preco FROM servicos WHERE ativo=1 ORDER BY nome"), ag=None)

@app.route('/agendamentos/editar/<int:id>', methods=['GET','POST'])
@login_required
def editar_agendamento(id):
    ag = query("SELECT * FROM agendamentos WHERE id=?", (id,), one=True)
    if not ag: flash('Não encontrado.','danger'); return redirect(url_for('agendamentos'))
    if request.method == 'POST':
        f = request.form
        query("UPDATE agendamentos SET cliente_id=?,barbeiro_id=?,servico_id=?,data=?,hora=?,status=?,obs=? WHERE id=?",
              (f['cliente_id'],f['barbeiro_id'],f['servico_id'],f['data'],f['hora'],f.get('status','Pendente'),f.get('obs',''),id), commit=True)
        flash('Agendamento atualizado!','success')
        return redirect(url_for('agendamentos'))
    return render_template('form_agendamento.html',
        clientes=query("SELECT id,nome FROM clientes ORDER BY nome"),
        barbeiros=query("SELECT id,nome FROM barbeiros WHERE ativo=1 ORDER BY nome"),
        servicos=query("SELECT id,nome,preco FROM servicos WHERE ativo=1 ORDER BY nome"), ag=dict(ag))

@app.route('/agendamentos/deletar/<int:id>', methods=['POST'])
@login_required
def deletar_agendamento(id):
    query("DELETE FROM agendamentos WHERE id=?",(id,), commit=True)
    flash('Agendamento excluído.','success')
    return redirect(url_for('agendamentos'))

# ── clientes ──────────────────────────────────────────────────────────────────
@app.route('/clientes')
@login_required
def clientes():
    rows = query("""SELECT c.*, (SELECT COUNT(*) FROM agendamentos WHERE cliente_id=c.id) visitas
                    FROM clientes c ORDER BY c.nome""")
    return render_template('admin_clientes.html', clientes=[dict(r) for r in rows])

@app.route('/clientes/novo', methods=['POST'])
@login_required
def novo_cliente():
    f = request.form
    query("INSERT INTO clientes(nome,telefone,email,bairro) VALUES(?,?,?,?)",
          (f['nome'],f.get('telefone',''),f.get('email',''),f.get('bairro','')), commit=True)
    flash('Cliente cadastrado!','success')
    return redirect(url_for('clientes'))

@app.route('/clientes/editar/<int:id>', methods=['GET','POST'])
@login_required
def editar_cliente(id):
    c = query("SELECT * FROM clientes WHERE id=?",(id,), one=True)
    if not c: flash('Não encontrado.','danger'); return redirect(url_for('clientes'))
    if request.method == 'POST':
        f = request.form
        query("UPDATE clientes SET nome=?,telefone=?,email=?,bairro=? WHERE id=?",
              (f['nome'],f.get('telefone',''),f.get('email',''),f.get('bairro',''),id), commit=True)
        flash('Cliente atualizado!','success')
        return redirect(url_for('clientes'))
    return render_template('admin_clientes.html', clientes=query("SELECT c.*, (SELECT COUNT(*) FROM agendamentos WHERE cliente_id=c.id) visitas FROM clientes c ORDER BY c.nome"), editar=dict(c))

@app.route('/clientes/deletar/<int:id>', methods=['POST'])
@login_required
def deletar_cliente(id):
    query("DELETE FROM clientes WHERE id=?",(id,), commit=True)
    flash('Cliente excluído.','success')
    return redirect(url_for('clientes'))

@app.route('/clientes/historico/<int:id>')
@login_required
def historico_cliente(id):
    return redirect(url_for('clientes'))  # placeholder

# ── serviços ──────────────────────────────────────────────────────────────────
@app.route('/servicos')
@login_required
def servicos():
    return render_template('admin_servicos.html', servicos=[dict(r) for r in query("SELECT * FROM servicos ORDER BY nome")])

@app.route('/servicos/novo', methods=['POST'])
@login_required
def novo_servico():
    f = request.form
    query("INSERT INTO servicos(nome,descricao,preco,duracao,ativo) VALUES(?,?,?,?,?)",
          (f['nome'],f.get('descricao',''),float(f.get('preco',0)),int(f.get('duracao',30)),int(f.get('ativo',1))), commit=True)
    flash('Serviço criado!','success')
    return redirect(url_for('servicos'))

@app.route('/servicos/editar/<int:id>', methods=['GET','POST'])
@login_required
def editar_servico(id):
    s = query("SELECT * FROM servicos WHERE id=?",(id,), one=True)
    if not s: flash('Não encontrado.','danger'); return redirect(url_for('servicos'))
    if request.method == 'POST':
        f = request.form
        query("UPDATE servicos SET nome=?,descricao=?,preco=?,duracao=?,ativo=? WHERE id=?",
              (f['nome'],f.get('descricao',''),float(f.get('preco',0)),int(f.get('duracao',30)),int(f.get('ativo',1)),id), commit=True)
        flash('Serviço atualizado!','success')
        return redirect(url_for('servicos'))
    return render_template('admin_servicos.html', servicos=query("SELECT * FROM servicos ORDER BY nome"), editar=dict(s))

@app.route('/servicos/deletar/<int:id>', methods=['POST'])
@login_required
def deletar_servico(id):
    query("DELETE FROM servicos WHERE id=?",(id,), commit=True)
    flash('Serviço excluído.','success')
    return redirect(url_for('servicos'))

# ── barbeiros ─────────────────────────────────────────────────────────────────
@app.route('/barbeiros')
@login_required
@admin_required
def barbeiros():
    rows = []
    for b in query("SELECT * FROM barbeiros ORDER BY nome"):
        b = dict(b)
        mes = date.today().strftime('%Y-%m')
        b['agendamentos_mes'] = query("SELECT COUNT(*) c FROM agendamentos WHERE barbeiro_id=? AND data LIKE ?",(b['id'],mes+'%'), one=True)['c']
        r = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.barbeiro_id=? AND a.data LIKE ? AND a.status='Concluido'",(b['id'],mes+'%'), one=True)['v']
        b['receita_mes'] = fmt_brl(r)
        rows.append(b)
    return render_template('admin_barbeiros.html', barbeiros=rows)

@app.route('/barbeiros/novo', methods=['POST'])
@login_required
@admin_required
def novo_barbeiro():
    f = request.form
    uid = None
    if f.get('usuario') and f.get('senha'):
        uid = query("INSERT INTO usuarios(nome,usuario,senha,role) VALUES(?,?,?,?)",
                    (f['nome'],f['usuario'],hash_pw(f['senha']),'barbeiro'), commit=True)
    query("INSERT INTO barbeiros(nome,telefone,especialidade,usuario_id,ativo) VALUES(?,?,?,?,?)",
          (f['nome'],f.get('telefone',''),f.get('especialidade',''),uid,int(f.get('ativo',1))), commit=True)
    flash('Barbeiro cadastrado!','success')
    return redirect(url_for('barbeiros'))

@app.route('/barbeiros/editar/<int:id>', methods=['GET','POST'])
@login_required
@admin_required
def editar_barbeiro(id):
    b = query("SELECT * FROM barbeiros WHERE id=?",(id,), one=True)
    if not b: flash('Não encontrado.','danger'); return redirect(url_for('barbeiros'))
    if request.method == 'POST':
        f = request.form
        query("UPDATE barbeiros SET nome=?,telefone=?,especialidade=?,ativo=? WHERE id=?",
              (f['nome'],f.get('telefone',''),f.get('especialidade',''),int(f.get('ativo',1)),id), commit=True)
        flash('Barbeiro atualizado!','success')
        return redirect(url_for('barbeiros'))
    return redirect(url_for('barbeiros'))

@app.route('/barbeiros/deletar/<int:id>', methods=['POST'])
@login_required
@admin_required
def deletar_barbeiro(id):
    query("DELETE FROM barbeiros WHERE id=?",(id,), commit=True)
    flash('Barbeiro removido.','success')
    return redirect(url_for('barbeiros'))

# ── relatórios ────────────────────────────────────────────────────────────────
@app.route('/relatorios')
@login_required
def relatorios():
    mes = date.today().strftime('%Y-%m')
    total_ag = query("SELECT COUNT(*) c FROM agendamentos WHERE data LIKE ?",(mes+'%',), one=True)['c']
    receita  = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.data LIKE ? AND a.status='Concluido'",(mes+'%',), one=True)['v']
    concl    = query("SELECT COUNT(*) c FROM agendamentos WHERE data LIKE ? AND status='Concluido'",(mes+'%',), one=True)['c']
    taxa     = round(concl/total_ag*100) if total_ag else 0
    ticket   = fmt_brl(receita/concl if concl else 0)
    # servicos populares
    sv = query("""SELECT s.nome, COUNT(*) total, SUM(s.preco) receita
                  FROM agendamentos a JOIN servicos s ON s.id=a.servico_id
                  WHERE a.data LIKE ? GROUP BY s.id ORDER BY total DESC LIMIT 8""",(mes+'%',))
    sv = [dict(r) for r in sv]
    mx = max((r['total'] for r in sv), default=1)
    for r in sv:
        r['pct'] = round(r['total']/mx*100)
        r['receita_num'] = float(r['receita'] or 0)
        r['receita'] = fmt_brl(r['receita'])
    # por barbeiro
    bp = query("""SELECT b.nome, COUNT(*) total, COALESCE(SUM(s.preco),0) receita
                  FROM agendamentos a JOIN barbeiros b ON b.id=a.barbeiro_id JOIN servicos s ON s.id=a.servico_id
                  WHERE a.data LIKE ? AND a.status='Concluido' GROUP BY b.id ORDER BY receita DESC""",(mes+'%',))
    bp = [dict(r) for r in bp]
    mx2 = max((r['receita'] for r in bp), default=1)
    for r in bp: r['pct'] = round(r['receita']/mx2*100); r['receita_num'] = float(r['receita']); r['receita'] = fmt_brl(r['receita'])
    # dias do mês com mais agendamentos (para gráfico de colunas)
    dias_mes = query("""SELECT CAST(SUBSTR(a.data,9,2) AS INTEGER) dia, COUNT(*) total
                        FROM agendamentos a WHERE a.data LIKE ?
                        GROUP BY dia ORDER BY dia""",(mes+'%',))
    dias_mes = [dict(r) for r in dias_mes]
    # clientes que mais agendaram (todo período)
    top_cli = query("""SELECT c.nome, COUNT(*) total, COALESCE(SUM(s.preco),0) gasto
                       FROM agendamentos a JOIN clientes c ON c.id=a.cliente_id
                       JOIN servicos s ON s.id=a.servico_id
                       WHERE a.status='Concluido'
                       GROUP BY c.id ORDER BY total DESC LIMIT 10""")
    top_cli = [dict(r) for r in top_cli]
    for r in top_cli: r['gasto'] = fmt_brl(r['gasto'])
    # ultimos agendamentos
    ultimos = query("""SELECT c.nome cliente, s.nome servico, b.nome barbeiro, a.data, a.hora, s.preco valor, a.status, a.pagamento
                       FROM agendamentos a JOIN clientes c ON c.id=a.cliente_id
                       JOIN servicos s ON s.id=a.servico_id JOIN barbeiros b ON b.id=a.barbeiro_id
                       ORDER BY a.data DESC, a.hora DESC LIMIT 20""")
    ul = [dict(r) for r in ultimos]
    for r in ul: r['valor'] = fmt_brl(r['valor'])
    return render_template('admin_relatorios.html',
        total_agendamentos=total_ag, receita_total=fmt_brl(receita),
        taxa_conclusao=taxa, ticket_medio=ticket,
        servicos_pop=sv, barbeiros_perf=bp,
        dias_mes=dias_mes, top_clientes=top_cli, ultimos=ul,
        mes_atual=date.today().strftime('%B/%Y').capitalize())

@app.route('/relatorios/exportar')
@login_required
def relatorios_exportar():
    """Gera planilha Excel completa com relatório do mês."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                                  GradientFill)
    from openpyxl.utils import get_column_letter

    mes = request.args.get('mes', date.today().strftime('%Y-%m'))
    mes_label = mes

    wb = Workbook()

    # ── helpers de estilo ──────────────────────────────────────────────
    PURPLE = '7C3AED'; PURPLE_LIGHT = 'EDE9FE'; AMBER = 'FBBF24'
    DARK   = '18181B'; GRAY    = '71717A'; LIGHT  = 'F4F4F5'
    SUCCESS_BG = 'D1FAE5'; SUCCESS_FG = '065F46'
    DANGER_BG  = 'FEE2E2'; DANGER_FG  = 'B91C1C'
    WARN_BG    = 'FEF9C3'; WARN_FG    = '854D0E'

    hdr_fill  = PatternFill('solid', fgColor=PURPLE)
    hdr_font  = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center')
    title_font = Font(bold=True, color=PURPLE, size=13)
    sub_font   = Font(color=GRAY, size=9)
    val_font   = Font(bold=True, color=DARK, size=10)
    thin = Side(style='thin', color='E4E4E7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def h(ws, row, col, value, bold=False, color=DARK, bg=None, size=10, align='left'):
        c = ws.cell(row=row, column=col, value=value)
        c.font = Font(bold=bold, color=color, size=size)
        c.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
        if bg: c.fill = PatternFill('solid', fgColor=bg)
        c.border = border
        return c

    def write_header_row(ws, row, cols):
        for ci, txt in enumerate(cols, 1):
            c = ws.cell(row=row, column=ci, value=txt)
            c.font = hdr_font; c.fill = hdr_fill
            c.alignment = hdr_align; c.border = border

    def auto_width(ws, extra=4):
        for col in ws.columns:
            w = max((len(str(c.value or '')) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + extra, 40)

    def section_title(ws, row, txt, span=6):
        ws.row_dimensions[row].height = 22
        c = ws.cell(row=row, column=1, value=txt)
        c.font = title_font
        c.fill = PatternFill('solid', fgColor=PURPLE_LIGHT)
        c.alignment = Alignment(horizontal='left', vertical='center')
        c.border = border
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)

    # ── ABA 1: Resumo ──────────────────────────────────────────────────
    ws1 = wb.active; ws1.title = 'Resumo'
    ws1.row_dimensions[1].height = 30
    c = ws1.cell(row=1, column=1, value=f'BarberBook PRO · Relatório {mes_label}')
    c.font = Font(bold=True, color=PURPLE, size=14)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws1.merge_cells('A1:F1')

    total_ag = query("SELECT COUNT(*) c FROM agendamentos WHERE data LIKE ?",(mes+'%',), one=True)['c']
    receita  = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.data LIKE ? AND a.status='Concluido'",(mes+'%',), one=True)['v']
    concl    = query("SELECT COUNT(*) c FROM agendamentos WHERE data LIKE ? AND status='Concluido'",(mes+'%',), one=True)['c']
    cancel   = query("SELECT COUNT(*) c FROM agendamentos WHERE data LIKE ? AND status='Cancelado'",(mes+'%',), one=True)['c']
    ticket   = round(float(receita)/concl,2) if concl else 0

    metrics = [
        ('Total de Agendamentos', total_ag),
        ('Agendamentos Concluídos', concl),
        ('Agendamentos Cancelados', cancel),
        ('Receita Total (R$)', f"R$ {fmt_brl(receita)}"),
        ('Ticket Médio (R$)', f"R$ {fmt_brl(ticket)}"),
        ('Taxa de Conclusão', f"{round(concl/total_ag*100) if total_ag else 0}%"),
    ]
    row = 3
    section_title(ws1, row, 'Indicadores do Mes', 2); row += 1
    for label, val in metrics:
        h(ws1, row, 1, label, bold=True, bg=LIGHT)
        h(ws1, row, 2, val, bold=True, color=PURPLE, align='center')
        row += 1

    # ── ABA 2: Agendamentos ────────────────────────────────────────────
    ws2 = wb.create_sheet('Agendamentos')
    c2 = ws2.cell(row=1, column=1, value=f'Agendamentos – {mes_label}')
    c2.font = Font(bold=True, color=PURPLE, size=13)
    ws2.merge_cells('A1:H1')
    write_header_row(ws2, 2, ['Data','Hora','Cliente','Serviço','Barbeiro','Valor (R$)','Pagamento','Status'])
    rows_ag = query("""SELECT a.data, a.hora, c.nome, s.nome, b.nome, s.preco, COALESCE(a.pagamento,'—'), a.status
                       FROM agendamentos a
                       JOIN clientes c ON c.id=a.cliente_id
                       JOIN servicos s ON s.id=a.servico_id
                       JOIN barbeiros b ON b.id=a.barbeiro_id
                       WHERE a.data LIKE ? ORDER BY a.data, a.hora""",(mes+'%',))
    status_styles = {'Concluido':(SUCCESS_BG,SUCCESS_FG),'Cancelado':(DANGER_BG,DANGER_FG),'Pendente':(WARN_BG,WARN_FG),'Confirmado':('EEF2FF','4338CA')}
    for ri, r in enumerate(rows_ag, 3):
        row_bg = LIGHT if ri % 2 == 0 else 'FFFFFF'
        for ci, val in enumerate(r, 1):
            if ci == 8:
                st = val; sbg, sfg = status_styles.get(st, ('F4F4F5','52525B'))
                h(ws2, ri, ci, val, bold=True, color=sfg, bg=sbg, align='center')
            elif ci == 6:
                h(ws2, ri, ci, float(val), bold=True, color=SUCCESS_FG, bg=row_bg, align='right')
            else:
                h(ws2, ri, ci, val, bg=row_bg)
    auto_width(ws2)
    ws2.freeze_panes = 'A3'

    # ── ABA 3: Serviços ────────────────────────────────────────────────
    ws3 = wb.create_sheet('Serviços')
    c3 = ws3.cell(row=1, column=1, value='Serviços Mais Solicitados'); c3.font = title_font
    ws3.merge_cells('A1:E1')
    write_header_row(ws3, 2, ['Serviço','Qtd Solicitações','Qtd Concluídos','Receita (R$)','% do Total'])
    svs = query("""SELECT s.nome,
                          COUNT(*) total,
                          SUM(CASE WHEN a.status='Concluido' THEN 1 ELSE 0 END) concl,
                          COALESCE(SUM(CASE WHEN a.status='Concluido' THEN s.preco ELSE 0 END),0) receita
                   FROM agendamentos a JOIN servicos s ON s.id=a.servico_id
                   WHERE a.data LIKE ? GROUP BY s.id ORDER BY total DESC""",(mes+'%',))
    svs = list(svs)
    tot_geral = sum(r[1] for r in svs) or 1
    for ri, r in enumerate(svs, 3):
        bg = LIGHT if ri % 2 == 0 else 'FFFFFF'
        pct = round(r[1]/tot_geral*100, 1)
        for ci, val in enumerate([r[0], r[1], r[2], f"R$ {fmt_brl(r[3])}", f"{pct}%"], 1):
            aln = 'right' if ci in (2,3,4,5) else 'left'
            h(ws3, ri, ci, val, bg=bg, align=aln)
    auto_width(ws3)

    # ── ABA 4: Clientes ────────────────────────────────────────────────
    ws4 = wb.create_sheet('Clientes')
    c4 = ws4.cell(row=1, column=1, value='Clientes – Engajamento'); c4.font = title_font
    ws4.merge_cells('A1:F1')
    write_header_row(ws4, 2, ['Cliente','Telefone','Total Visitas','Concluídos','Cancelados','Total Gasto (R$)'])
    clis = query("""SELECT c.nome, COALESCE(c.telefone,'—'),
                           COUNT(*) total,
                           SUM(CASE WHEN a.status='Concluido' THEN 1 ELSE 0 END) concl,
                           SUM(CASE WHEN a.status='Cancelado' THEN 1 ELSE 0 END) canc,
                           COALESCE(SUM(CASE WHEN a.status='Concluido' THEN s.preco ELSE 0 END),0) gasto
                    FROM agendamentos a JOIN clientes c ON c.id=a.cliente_id
                    JOIN servicos s ON s.id=a.servico_id
                    GROUP BY c.id ORDER BY total DESC""")
    for ri, r in enumerate(clis, 3):
        bg = LIGHT if ri % 2 == 0 else 'FFFFFF'
        for ci, val in enumerate([r[0], r[1], r[2], r[3], r[4], f"R$ {fmt_brl(r[5])}"], 1):
            aln = 'right' if ci >= 3 else 'left'
            h(ws4, ri, ci, val, bg=bg, align=aln)
    auto_width(ws4)

    # ── ABA 5: Barbeiros ───────────────────────────────────────────────
    ws5 = wb.create_sheet('Barbeiros')
    c5 = ws5.cell(row=1, column=1, value='Performance por Barbeiro'); c5.font = title_font
    ws5.merge_cells('A1:G1')
    write_header_row(ws5, 2, ['Barbeiro','Total Atend.','Concluídos','Cancelados','Receita (R$)','Ticket Médio','% Conclusão'])
    barbs = query("""SELECT b.nome,
                            COUNT(*) total,
                            SUM(CASE WHEN a.status='Concluido' THEN 1 ELSE 0 END) concl,
                            SUM(CASE WHEN a.status='Cancelado' THEN 1 ELSE 0 END) canc,
                            COALESCE(SUM(CASE WHEN a.status='Concluido' THEN s.preco ELSE 0 END),0) rec
                     FROM agendamentos a JOIN barbeiros b ON b.id=a.barbeiro_id
                     JOIN servicos s ON s.id=a.servico_id
                     GROUP BY b.id ORDER BY rec DESC""")
    for ri, r in enumerate(barbs, 3):
        bg = LIGHT if ri % 2 == 0 else 'FFFFFF'
        tc = f"R$ {fmt_brl(r[4]/r[2] if r[2] else 0)}"
        pct = f"{round(r[2]/r[1]*100) if r[1] else 0}%"
        for ci, val in enumerate([r[0], r[1], r[2], r[3], f"R$ {fmt_brl(r[4])}", tc, pct], 1):
            aln = 'right' if ci >= 2 else 'left'
            h(ws5, ri, ci, val, bg=bg, align=aln)
    auto_width(ws5)

    # ── output ─────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    from flask import send_file
    fname = f"barberbook_relatorio_{mes}.xlsx"
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=fname)

# ── financeiro ────────────────────────────────────────────────────────────────
@app.route('/financeiro')
@login_required
@admin_required
def financeiro():
    cfg = {r['chave']:r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    hoje = date.today().isoformat(); mes = date.today().strftime('%Y-%m')
    sem_ini = (date.today() - timedelta(days=date.today().weekday())).isoformat()
    r_hoje = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.data=? AND a.status='Concluido'",(hoje,), one=True)['v']
    r_sem  = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.data>=? AND a.status='Concluido'",(sem_ini,), one=True)['v']
    r_mes  = query("SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id WHERE a.data LIKE ? AND a.status='Concluido'",(mes+'%',), one=True)['v']
    formas = [{'forma':'Pix','pct':60,'total':fmt_brl(float(r_mes)*0.6)},
              {'forma':'Cartão Débito','pct':25,'total':fmt_brl(float(r_mes)*0.25)},
              {'forma':'Dinheiro','pct':15,'total':fmt_brl(float(r_mes)*0.15)}]
    receitas = query("""SELECT a.data, c.nome cliente, s.nome servico, b.nome barbeiro, s.preco valor
                        FROM agendamentos a JOIN clientes c ON c.id=a.cliente_id
                        JOIN servicos s ON s.id=a.servico_id JOIN barbeiros b ON b.id=a.barbeiro_id
                        WHERE a.status='Concluido' ORDER BY a.data DESC LIMIT 15""")
    rc = [dict(r) for r in receitas]
    for r in rc: r['valor'] = fmt_brl(r['valor']); r['forma'] = 'Pix'
    return render_template('admin_financeiro.html',
        receita_hoje=fmt_brl(r_hoje), receita_semana=fmt_brl(r_sem), receita_mes=fmt_brl(r_mes),
        formas_pagamento=formas, receitas=rc,
        vencimento=cfg.get('vencimento',''), vencimento_iso=cfg.get('vencimento_iso',''))

# ── contratos ──────────────────────────────────────────────────────────────────
@app.route('/contratos')
@login_required
@admin_required
def contratos():
    if session.get('demo_mode'):
        return redirect(url_for('dashboard'))
    cfg = {r['chave']:r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    contratos = [{'barbearia':cfg.get('nome','Barbearia'), 'plano':'PRO Mensal',
                  'inicio':'01/01/2025', 'vencimento':cfg.get('vencimento',''),
                  'vencimento_iso':cfg.get('vencimento_iso',''), 'status':'Ativo'}]
    return render_template('admin_contratos.html', contratos=contratos)

@app.route('/contratos/pdf')
@login_required
def contrato_pdf():
    flash('Funcionalidade de PDF em desenvolvimento.','info')
    return redirect(url_for('contratos'))

# ── configurações ──────────────────────────────────────────────────────────────
@app.route('/configuracoes')
@login_required
@admin_required
def configuracoes():
    cfg = {r['chave']:r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    dias_semana = []
    for k,n in [('seg','Segunda'),('ter','Terça'),('qua','Quarta'),('qui','Quinta'),('sex','Sexta'),('sab','Sábado'),('dom','Domingo')]:
        h = query("SELECT * FROM horarios WHERE dia=?",(k,), one=True)
        if h: dias_semana.append({'key':k,'nome':n,'ativo':h['ativo'],'inicio':h['inicio'],'fim':h['fim']})
        else: dias_semana.append({'key':k,'nome':n,'ativo':1,'inicio':'08:00','fim':'18:00'})
    return render_template('admin_configuracoes.html', config=cfg, dias_semana=dias_semana)

@app.route('/configuracoes/salvar', methods=['POST'])
@login_required
@admin_required
def salvar_config():
    for k in ['nome','telefone','email','endereco','instagram','link_agendamento',
              'whatsapp_numero','whatsapp_phone_id','whatsapp_api_token','whatsapp_verify_token']:
        query("INSERT OR REPLACE INTO configuracoes(chave,valor) VALUES(?,?)",(k, request.form.get(k,'')), commit=True)
    flash('Configurações salvas!','success')
    return redirect(url_for('configuracoes'))

@app.route('/configuracoes/horarios', methods=['POST'])
@login_required
@admin_required
def salvar_horarios():
    for k in ['seg','ter','qua','qui','sex','sab','dom']:
        ativo = 1 if request.form.get(f'ativo_{k}') else 0
        inicio = request.form.get(f'inicio_{k}','08:00')
        fim    = request.form.get(f'fim_{k}','18:00')
        query("INSERT OR REPLACE INTO horarios(dia,ativo,inicio,fim) VALUES(?,?,?,?)",(k,ativo,inicio,fim), commit=True)
    flash('Horários salvos!','success')
    return redirect(url_for('configuracoes'))

@app.route('/configuracoes/senha', methods=['POST'])
@login_required
def alterar_senha():
    uid = session['uid']
    u   = query("SELECT senha FROM usuarios WHERE id=?",(uid,), one=True)
    f   = request.form
    if not u or u['senha'] != hash_pw(f.get('senha_atual','')):
        flash('Senha atual incorreta.','danger')
    elif f.get('nova_senha') != f.get('confirmar'):
        flash('Senhas não coincidem.','danger')
    else:
        query("UPDATE usuarios SET senha=? WHERE id=?",(hash_pw(f['nova_senha']),uid), commit=True)
        flash('Senha alterada com sucesso!','success')
    return redirect(url_for('configuracoes'))

# ── perfil ─────────────────────────────────────────────────────────────────────
@app.route('/perfil')
@login_required
def perfil():
    return render_template('perfil.html')

@app.route('/perfil/salvar', methods=['POST'])
@login_required
def salvar_perfil():
    uid = session['uid']
    f   = request.form
    query("UPDATE usuarios SET nome=?,usuario=? WHERE id=?",(f['nome'],f['usuario'],uid), commit=True)
    if f.get('nova_senha') and f['nova_senha'] == f.get('confirmar'):
        query("UPDATE usuarios SET senha=? WHERE id=?",(hash_pw(f['nova_senha']),uid), commit=True)
    session['nome'] = f['nome']; session['usuario'] = f['usuario']
    flash('Perfil atualizado!','success')
    return redirect(url_for('perfil'))

@app.route('/whatsapp')
@login_required
@admin_required
def whatsapp_central():
    mensagens = [dict(r) for r in query("""
        SELECT w.*, a.data ag_data, a.hora ag_hora
        FROM whatsapp_mensagens w
        LEFT JOIN agendamentos a ON a.id=w.agendamento_id
        ORDER BY w.id DESC LIMIT 80
    """)]
    pendencias = [dict(r) for r in query("""
        SELECT a.id, a.data, a.hora, a.status, c.nome cliente, b.nome barbeiro, s.nome servico
        FROM agendamentos a
        JOIN clientes c ON c.id=a.cliente_id
        JOIN barbeiros b ON b.id=a.barbeiro_id
        JOIN servicos s ON s.id=a.servico_id
        WHERE a.status='Solicitado via WhatsApp'
        ORDER BY a.data, a.hora LIMIT 40
    """)]
    return render_template(
        'admin_whatsapp.html',
        mensagens=mensagens,
        pendencias=pendencias,
        webhook_url=f"{request.url_root.rstrip('/')}{url_for('whatsapp_webhook')}"
    )

@app.route('/whatsapp/marcar-lidas', methods=['POST'])
@login_required
@admin_required
def whatsapp_marcar_lidas():
    query("UPDATE whatsapp_mensagens SET lida=1 WHERE lida=0", commit=True)
    flash('Notificações do WhatsApp marcadas como lidas.', 'success')
    return redirect(url_for('whatsapp_central'))

@app.route('/api/whatsapp/unread')
@login_required
def api_whatsapp_unread():
    if session.get('role') == 'barbeiro':
        return jsonify({'count': 0})
    unread = query("SELECT COUNT(*) c FROM whatsapp_mensagens WHERE lida=0", one=True)['c']
    pend = query("SELECT COUNT(*) c FROM agendamentos WHERE status='Solicitado via WhatsApp'", one=True)['c']
    return jsonify({'count': int(unread or 0) + int(pend or 0)})

@app.route('/webhook/whatsapp', methods=['GET', 'POST'])
def whatsapp_webhook():
    if request.method == 'GET':
        cfg = _cfg()
        verify_token = (cfg.get('whatsapp_verify_token') or '').strip()
        mode = request.args.get('hub.mode')
        token = request.args.get('hub.verify_token')
        challenge = request.args.get('hub.challenge', '')
        if mode == 'subscribe' and verify_token and token == verify_token:
            return challenge, 200
        return 'forbidden', 403

    if request.form.get('Body'):
        _registrar_whatsapp_entrada(
            request.form.get('ProfileName') or 'Cliente',
            request.form.get('From', ''),
            request.form.get('Body', ''),
            origem='twilio',
            payload=request.form.to_dict()
        )
        return 'ok', 200

    data = request.get_json(silent=True) or {}
    try:
        for entry in data.get('entry', []):
            for change in entry.get('changes', []):
                value = change.get('value', {}) or {}
                contatos = value.get('contacts') or []
                nome = 'Cliente'
                if contatos:
                    nome = contatos[0].get('profile', {}).get('name', 'Cliente')
                for msg in value.get('messages', []) or []:
                    texto = (msg.get('text') or {}).get('body', '')
                    if texto:
                        _registrar_whatsapp_entrada(nome, msg.get('from', ''), texto, origem='meta', payload=msg)
        return jsonify({'ok': True})
    except Exception as exc:
        return jsonify({'ok': False, 'msg': str(exc)}), 200

# ── painel do barbeiro ────────────────────────────────────────────────────────
import urllib.parse

def _get_barbeiro():
    """Retorna o registro de barbeiros ligado ao usuário logado."""
    return query("SELECT * FROM barbeiros WHERE usuario_id=?", (session['uid'],), one=True)

def _wpp_link(ag_id, msg_template):
    """Gera link WhatsApp para o cliente do agendamento."""
    c = query("""SELECT cl.nome, cl.telefone, a.data, a.hora,
                        COALESCE(s.nome,'Serviço') servico,
                        COALESCE(b.nome,'Barbeiro') barbeiro
                 FROM agendamentos a
                 JOIN clientes cl ON a.cliente_id=cl.id
                 LEFT JOIN servicos s ON s.id=a.servico_id
                 LEFT JOIN barbeiros b ON b.id=a.barbeiro_id
                 WHERE a.id=?""", (ag_id,), one=True)
    if not c or not c['telefone']:
        return ''
    tel = ''.join(filter(str.isdigit, c['telefone']))
    if len(tel) < 10:
        return ''
    repl = {
        '{nome}': (c['nome'] or 'Cliente').split()[0],
        '{servico}': c['servico'] or 'Serviço',
        '{barbeiro}': c['barbeiro'] or 'Barbeiro',
        '{data}': c['data'] or '',
        '{hora}': c['hora'] or ''
    }
    msg = msg_template
    for k, v in repl.items():
        msg = msg.replace(k, str(v))
    tel = tel if tel.startswith('55') else f'55{tel}'
    return f"https://wa.me/{tel}?text={urllib.parse.quote(msg)}"

def _notificar_cliente_agendamento(ag_id, msg_template):
    link = _wpp_link(ag_id, msg_template)
    c = query("""SELECT cl.telefone FROM clientes cl
                 JOIN agendamentos a ON a.cliente_id=cl.id WHERE a.id=?""", (ag_id,), one=True)
    if c and c['telefone']:
        repl_msg = urllib.parse.unquote(link.split('text=', 1)[1]) if 'text=' in link else ''
        envio = _send_whatsapp_message(c['telefone'], repl_msg)
    else:
        envio = {'ok': False, 'msg': 'Cliente sem telefone.'}
    return {'link': link, 'envio': envio}

@app.route('/barbeiro')
@login_required
def barbeiro_painel():
    if session.get('role') in ('admin', 'gerente'):
        return redirect(url_for('dashboard'))
    b = _get_barbeiro()
    if not b:
        flash('Seu usuário não está vinculado a nenhum barbeiro. Fale com o admin.', 'warning')
        return redirect(url_for('logout'))
    hoje = date.today().isoformat()
    fim_semana = (date.today() + timedelta(days=6)).isoformat()
    mes = date.today().strftime('%Y-%m')
    def load(extra, args=()):
        rows = query(f"""
            SELECT a.id, a.hora, a.data, a.status, a.pagamento, a.obs,
                   c.nome cliente, c.telefone telefone,
                   s.nome servico, s.preco preco, s.duracao duracao
            FROM agendamentos a
            JOIN clientes c ON c.id=a.cliente_id
            JOIN servicos s ON s.id=a.servico_id
            WHERE a.barbeiro_id=? {extra}
            ORDER BY a.data, a.hora
        """, (b['id'],) + args)
        result = [dict(r) for r in rows]
        for r in result: r['preco'] = fmt_brl(r['preco'])
        return result
    ag_hoje   = load("AND a.data=?", (hoje,))
    ag_semana = load("AND a.data>=? AND a.data<=?", (hoje, fim_semana))
    ag_mes    = load("AND a.data LIKE ?", (mes+'%',))
    receita_hoje = query("""SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id
        WHERE a.barbeiro_id=? AND a.data=? AND a.status='Concluido'""", (b['id'], hoje), one=True)['v']
    receita_mes = query("""SELECT COALESCE(SUM(s.preco),0) v FROM agendamentos a JOIN servicos s ON s.id=a.servico_id
        WHERE a.barbeiro_id=? AND a.data LIKE ? AND a.status='Concluido'""", (b['id'], mes+'%'), one=True)['v']
    cfg = {r['chave']:r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    return render_template('barbeiro_painel.html',
        b=dict(b), ag_hoje=ag_hoje, ag_semana=ag_semana, ag_mes=ag_mes,
        hoje_fmt=date.today().strftime('%d/%m/%Y'),
        receita_hoje=fmt_brl(receita_hoje), receita_mes=fmt_brl(receita_mes),
        barbearia_nome=cfg.get('nome',''))

@app.route('/barbeiro/status', methods=['POST'])
@login_required
def barbeiro_toggle_status():
    b = _get_barbeiro()
    if not b: return jsonify({'ok': False}), 403
    novo = 0 if b['em_atendimento'] else 1
    query("UPDATE barbeiros SET em_atendimento=? WHERE id=?", (novo, b['id']), commit=True)
    return jsonify({'ok': True, 'em_atendimento': bool(novo)})

@app.route('/barbeiro/confirmar/<int:id>', methods=['POST'])
@login_required
def barbeiro_confirmar(id):
    b = _get_barbeiro()
    ag = query("SELECT * FROM agendamentos WHERE id=?", (id,), one=True)
    if not ag: return jsonify({'ok': False, 'msg': 'Não encontrado'}), 404
    if b and ag['barbeiro_id'] != b['id'] and session.get('role') not in ('admin','gerente'):
        return jsonify({'ok': False, 'msg': 'Sem permissão'}), 403
    pag = request.form.get('pagamento', '')
    query("UPDATE agendamentos SET status='Concluido', pagamento=? WHERE id=?", (pag, id), commit=True)
    notif = _notificar_cliente_agendamento(id, 'Olá {nome}! Seu atendimento de *{servico}* foi *concluido* com sucesso. Obrigado pela visita!')
    return jsonify({'ok': True, 'wpp': notif['link'], 'pagamento': pag, 'sent': notif['envio']['ok']})

@app.route('/barbeiro/iniciar/<int:id>', methods=['POST'])
@login_required
def barbeiro_iniciar(id):
    b = _get_barbeiro()
    ag = query("SELECT * FROM agendamentos WHERE id=?", (id,), one=True)
    if not ag: return jsonify({'ok': False}), 404
    if b and ag['barbeiro_id'] != b['id'] and session.get('role') not in ('admin','gerente'):
        return jsonify({'ok': False}), 403
    query("UPDATE agendamentos SET status='Confirmado' WHERE id=?", (id,), commit=True)
    notif = _notificar_cliente_agendamento(id, 'Olá {nome}! Seu agendamento de *{servico}* foi *confirmado* para {data} às {hora} com {barbeiro}. Aguardamos voce!')
    return jsonify({'ok': True, 'wpp': notif['link'], 'sent': notif['envio']['ok']})

@app.route('/barbeiro/cancelar/<int:id>', methods=['POST'])
@login_required
def barbeiro_cancelar(id):
    b = _get_barbeiro()
    ag = query("SELECT * FROM agendamentos WHERE id=?", (id,), one=True)
    if not ag: return jsonify({'ok': False, 'msg': 'Não encontrado'}), 404
    if b and ag['barbeiro_id'] != b['id'] and session.get('role') not in ('admin','gerente'):
        return jsonify({'ok': False, 'msg': 'Sem permissão'}), 403
    motivo = request.form.get('motivo', '').strip()
    query("UPDATE agendamentos SET status='Cancelado' WHERE id=?", (id,), commit=True)
    mot_txt = f' Motivo: _{motivo}_.' if motivo else ''
    notif = _notificar_cliente_agendamento(id, f'Olá {{nome}}! Infelizmente seu agendamento foi *cancelado*.{mot_txt} Entre em contato para reagendar.')
    return jsonify({'ok': True, 'wpp': notif['link'], 'sent': notif['envio']['ok']})

@app.route('/barbeiro/config', methods=['GET', 'POST'])
@login_required
def barbeiro_config():
    if session.get('role') in ('admin', 'gerente'):
        return redirect(url_for('configuracoes'))
    uid = session['uid']
    u = query("SELECT * FROM usuarios WHERE id=?", (uid,), one=True)
    b = _get_barbeiro()
    if request.method == 'POST':
        f = request.form
        nome     = f.get('nome', '').strip()
        email    = f.get('email', '').strip()
        usuario  = f.get('usuario', '').strip()
        # foto: salva nome do arquivo enviado (base64 não recomendado – apenas filename)
        foto = u['foto'] if u else ''
        if 'foto' in request.files:
            arq = request.files['foto']
            if arq and arq.filename:
                import os as _os
                ext = _os.path.splitext(arq.filename)[1].lower()
                if ext in ('.jpg','.jpeg','.png','.webp'):
                    pasta = _os.path.join(_os.path.dirname(__file__), 'static', 'img', 'perfis')
                    _os.makedirs(pasta, exist_ok=True)
                    fname = f"barb_{uid}{ext}"
                    arq.save(_os.path.join(pasta, fname))
                    foto = f"img/perfis/{fname}"
        query("UPDATE usuarios SET nome=?,email=?,usuario=?,foto=? WHERE id=?",
              (nome, email, usuario, foto, uid), commit=True)
        if b:
            query("UPDATE barbeiros SET nome=?,email=? WHERE id=?", (nome, email, b['id']), commit=True)
        if f.get('nova_senha'):
            if not u or u['senha'] != hash_pw(f.get('senha_atual', '')):
                flash('Senha atual incorreta.', 'danger')
            elif f['nova_senha'] != f.get('confirmar', ''):
                flash('Senhas não coincidem.', 'danger')
            else:
                query("UPDATE usuarios SET senha=? WHERE id=?", (hash_pw(f['nova_senha']), uid), commit=True)
                flash('Senha alterada!', 'success')
        session['nome'] = nome
        session['usuario'] = usuario
        session['email'] = email
        flash('Perfil atualizado!', 'success')
        return redirect(url_for('barbeiro_config'))
    return render_template('barbeiro_config.html',
        u=dict(u) if u else {}, b=dict(b) if b else {})

# ── web push helpers ─────────────────────────────────────────────────────────
def _get_vapid_pem():
    row = query("SELECT valor FROM configuracoes WHERE chave='vapid_private'", one=True)
    return row['valor'] if row else None

def _notify_barber_async(barbeiro_id, title, body):
    """Envia push notification ao barbeiro de forma assíncrona."""
    if not _PUSH_OK:
        return
    def _send():
        with app.app_context():
            pem = _get_vapid_pem()
            if not pem:
                return
            subs = query("SELECT id, subscription_json FROM push_subscriptions WHERE barbeiro_id=?", (barbeiro_id,))
            payload = json.dumps({"title": title, "body": body, "icon": "/static/img/icon-192.png"})
            to_delete = []
            for sub in subs:
                try:
                    _webpush(
                        json.loads(sub['subscription_json']),
                        data=payload,
                        vapid_private_key=pem,
                        vapid_claims=_VAPID_CLAIMS
                    )
                except WebPushException as ex:
                    if ex.response and ex.response.status_code in (404, 410):
                        to_delete.append(sub['id'])
                except Exception:
                    pass
            for did in to_delete:
                query("DELETE FROM push_subscriptions WHERE id=?", (did,), commit=True)
    threading.Thread(target=_send, daemon=True).start()

@app.route('/sw.js')
def service_worker():
    """Serve o service worker na raiz para permitir escopo global."""
    from flask import send_from_directory
    return send_from_directory(
        os.path.join(os.path.dirname(__file__), 'static', 'js'),
        'sw.js',
        mimetype='application/javascript'
    )

@app.route('/api/vapid-key')
def api_vapid_key():
    row = query("SELECT valor FROM configuracoes WHERE chave='vapid_public'", one=True)
    return jsonify({'public_key': row['valor'] if row else None})

@app.route('/api/cliente-lookup')
def api_cliente_lookup():
    """Retorna últimos serviços do cliente pelo telefone para sugestão WPP."""
    tel = ''.join(filter(str.isdigit, request.args.get('telefone', '')))
    if len(tel) < 8:
        return jsonify({'found': False})
    cli = query("SELECT id, nome FROM clientes WHERE REPLACE(REPLACE(REPLACE(REPLACE(telefone,'(',''),')',''),' ',''),'-','') LIKE ?",
                ('%'+tel[-8:],), one=True)
    if not cli:
        return jsonify({'found': False})
    ultimos = query("""SELECT s.nome servico, s.id servico_id, b.nome barbeiro, b.id barbeiro_id,
                              a.data, a.hora
                       FROM agendamentos a
                       JOIN servicos s ON s.id=a.servico_id
                       JOIN barbeiros b ON b.id=a.barbeiro_id
                       WHERE a.cliente_id=? AND a.status='Concluido'
                       ORDER BY a.data DESC, a.hora DESC LIMIT 3""", (cli['id'],))
    ultimos = [dict(r) for r in ultimos]
    cfg = {r['chave']: r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    return jsonify({'found': True, 'nome': cli['nome'], 'ultimos': ultimos,
                    'barbearia': cfg.get('nome','Barbearia'), 'tel_barbearia': cfg.get('telefone','')})

@app.route('/barbeiro/push-registrar', methods=['POST'])
@login_required
def barbeiro_push_registrar():
    b = _get_barbeiro()
    if not b:
        return jsonify({'ok': False}), 403
    data = request.get_json(force=True)
    if not data:
        return jsonify({'ok': False}), 400
    sub_str = json.dumps(data)
    try:
        query("INSERT OR REPLACE INTO push_subscriptions(barbeiro_id,subscription_json) VALUES(?,?)",
              (b['id'], sub_str), commit=True)
    except Exception:
        pass
    return jsonify({'ok': True})

# ── landing page pública ──────────────────────────────────────────────────────
@app.route('/', methods=['GET'])
@app.route('/agendar', methods=['GET', 'POST'])
def landing():
    if request.method == 'POST':
        f = request.get_json(force=True) if request.is_json else request.form
        nome     = str(f.get('nome', '')).strip()
        telefone = str(f.get('telefone', '')).strip()
        servico_id  = f.get('servico_id')
        barbeiro_id = f.get('barbeiro_id')
        data_ag  = str(f.get('data', ''))
        hora_ag  = str(f.get('hora', ''))
        if not all([nome, telefone, servico_id, barbeiro_id, data_ag, hora_ag]):
            return jsonify({'ok': False, 'msg': 'Preencha todos os campos.'}), 400
        try:
            d = date.fromisoformat(data_ag)
            if d < date.today():
                raise ValueError('passado')
        except Exception:
            return jsonify({'ok': False, 'msg': 'Data inválida.'}), 400
        cli = query("SELECT id FROM clientes WHERE telefone=?", (telefone,), one=True)
        if cli:
            cli_id = cli['id']
            query("UPDATE clientes SET nome=? WHERE id=?", (nome, cli_id), commit=True)
        else:
            cli_id = query("INSERT INTO clientes(nome,telefone) VALUES(?,?)", (nome, telefone), commit=True)
        clash = query(
            "SELECT id FROM agendamentos WHERE data=? AND hora=? AND barbeiro_id=? AND status NOT IN ('Cancelado')",
            (data_ag, hora_ag, int(barbeiro_id))
        )
        if clash:
            return jsonify({'ok': False, 'msg': 'Horário não está mais disponível. Escolha outro.'}), 409
        ag_id = query(
            "INSERT INTO agendamentos(cliente_id,barbeiro_id,servico_id,data,hora,status) VALUES(?,?,?,?,?,'Pendente')",
            (cli_id, int(barbeiro_id), int(servico_id), data_ag, hora_ag), commit=True
        )
        srv_row = query("SELECT nome FROM servicos WHERE id=?", (servico_id,), one=True)
        barb_row = query("SELECT nome FROM barbeiros WHERE id=?", (barbeiro_id,), one=True)
        _notify_barber_async(
            int(barbeiro_id),
            '\U0001f4c5 Novo Agendamento!',
            f'{nome} \u00b7 {srv_row["nome"] if srv_row else ""} \u00b7 {hora_ag}'
        )
        notif = _notificar_cliente_agendamento(
            ag_id,
            'Olá {nome}! Seu agendamento de *{servico}* foi recebido para {data} às {hora} com {barbeiro}. Responda aqui para confirmar'
        )
        return jsonify({'ok': True, 'ag_id': ag_id,
                        'barbeiro': barb_row['nome'] if barb_row else '',
                        'whatsapp_enviado': notif['envio']['ok'],
                        'msg': 'Agendamento confirmado!'})
    # GET
    cfg = {r['chave']: r['valor'] for r in query("SELECT chave,valor FROM configuracoes")}
    servicos = [dict(r) for r in query("SELECT id,nome,descricao,preco,duracao FROM servicos WHERE ativo=1")]
    for s in servicos:
        s['preco_fmt'] = fmt_brl(s['preco'])
    barbeiros = [dict(r) for r in query("SELECT id,nome,especialidade FROM barbeiros WHERE ativo=1")]
    _dias_ordem = {'seg':1,'ter':2,'qua':3,'qui':4,'sex':5,'sab':6,'dom':7}
    _dias_nomes = {'seg':'Segunda','ter':'Terça','qua':'Quarta','qui':'Quinta','sex':'Sexta','sab':'Sábado','dom':'Domingo'}
    _dias_abrev = {'seg':'Seg','ter':'Ter','qua':'Qua','qui':'Qui','sex':'Sex','sab':'Sáb','dom':'Dom'}
    hrs_raw = query("SELECT dia,ativo,inicio,fim FROM horarios")
    horarios = sorted(
        [{'dia': r['dia'], 'nome': _dias_nomes.get(r['dia'], r['dia']),
          'abrev': _dias_abrev.get(r['dia'], r['dia']),
          'ativo': bool(r['ativo']),
          'inicio': r['inicio'] or '08:00', 'fim': r['fim'] or '18:00'}
         for r in hrs_raw],
        key=lambda x: _dias_ordem.get(x['dia'], 9)
    )
    return render_template('landing.html', cfg=cfg, servicos=servicos, barbeiros=barbeiros, horarios=horarios)

@app.route('/api/horarios-disponiveis')
def api_horarios():
    data_str    = request.args.get('data', '')
    barbeiro_id = request.args.get('barbeiro_id', 'qualquer')
    servico_id  = request.args.get('servico_id', '')
    try:
        d = date.fromisoformat(data_str)
    except Exception:
        return jsonify({'slots': [], 'fechado': True})
    if d < date.today():
        return jsonify({'slots': [], 'fechado': True, 'msg': 'Data no passado'})
    dias_map = {0: 'seg', 1: 'ter', 2: 'qua', 3: 'qui', 4: 'sex', 5: 'sab', 6: 'dom'}
    dia_key = dias_map[d.weekday()]
    h = query("SELECT ativo,inicio,fim FROM horarios WHERE dia=?", (dia_key,), one=True)
    if not h or not h['ativo']:
        return jsonify({'slots': [], 'fechado': True, 'msg': 'Fechado neste dia'})
    srv = query("SELECT duracao FROM servicos WHERE id=?", (servico_id,), one=True) if servico_id else None
    duracao = int(srv['duracao']) if srv else 30
    ini = datetime.strptime(h['inicio'], '%H:%M')
    fim = datetime.strptime(h['fim'], '%H:%M')
    all_slots = []
    t = ini
    while t + timedelta(minutes=duracao) <= fim:
        all_slots.append(t.strftime('%H:%M'))
        t += timedelta(minutes=30)
    if barbeiro_id == 'qualquer':
        all_barb_ids = [r['id'] for r in query("SELECT id FROM barbeiros WHERE ativo=1")]
    else:
        try:
            all_barb_ids = [int(barbeiro_id)]
        except Exception:
            all_barb_ids = []
    result = []
    for slot in all_slots:
        booked = set(r['barbeiro_id'] for r in query(
            "SELECT barbeiro_id FROM agendamentos WHERE data=? AND hora=? AND status NOT IN ('Cancelado')",
            (data_str, slot)
        ))
        free = [bid for bid in all_barb_ids if bid not in booked]
        if free:
            result.append({'hora': slot, 'barbeiro_id': free[0]})
    return jsonify({'slots': result, 'fechado': False})

# ── run ────────────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    init_db()
    port = int(os.environ.get('PORT', 5001))
    app.run(debug=False, host='0.0.0.0', port=port)
